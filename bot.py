import os
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from telegram import Update, ReplyKeyboardMarkup, ReplyKeyboardRemove
from telegram.ext import (
    Application,
    CommandHandler,
    ConversationHandler,
    MessageHandler,
    filters,
)

TYPE, INPUT_MODE, AMOUNT, BULK_TEXT, PAYMENT, FILENAME = range(6)

DATA_DIR = Path("data")
DATA_DIR.mkdir(exist_ok=True)

HEADERS = ["Turi", "Miqdor", "Izoh", "Naqd yoki Perechisleniya"]

HEADER_FONT = Font(bold=True, size=12)
HEADER_ALIGNMENT = Alignment(horizontal="center")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
COLUMN_WIDTHS = [15, 18, 30, 25]


def get_user_file(user_id: int) -> Path:
    return DATA_DIR / f"{user_id}.xlsx"


def ensure_workbook(user_id: int) -> Path:
    path = get_user_file(user_id)
    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Hisobot"
        for col_idx, header in enumerate(HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER
        for col_idx, width in enumerate(COLUMN_WIDTHS, 1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
        wb.save(path)
    return path


def evaluate_math(expression: str) -> float | None:
    cleaned = expression.replace(" ", "").replace(",", ".")
    if not re.match(r"^[\d.+\-*/()]+$", cleaned):
        return None
    try:
        result = eval(cleaned, {"__builtins__": {}}, {})  # noqa: S307
        return float(result)
    except Exception:
        return None


def _parse_bulk_amount(text: str) -> tuple[float, str] | None:
    m = re.match(r"^(\$?)(\d{1,3}(?:\.\d{3})+)(\$?)\s*(.*)$", text)
    if m:
        dollar = m.group(1) or m.group(3)
        amount = float(m.group(2).replace(".", ""))
        if dollar:
            amount *= 12500
        return amount, (m.group(4) or "").strip()
    m = re.match(r"^\$(\d+)\s*(.*)$", text)
    if m:
        return float(m.group(1)) * 12500, (m.group(2) or "").strip()
    m = re.match(r"^(\d+)\$\s*(.*)$", text)
    if m:
        return float(m.group(1)) * 12500, (m.group(2) or "").strip()
    return None


def parse_bulk_text(text: str) -> list[tuple[float, str]]:
    lines = text.strip().split("\n")

    merged = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("+") and merged:
            merged[-1] += line
        else:
            merged.append(line)

    results = []
    for line in merged:
        line = line.strip().strip("+").strip()
        if not line:
            continue

        segments = [s.strip() for s in line.split("+") if s.strip()]
        total = 0.0
        descs = []
        for seg in segments:
            parsed = _parse_bulk_amount(seg)
            if parsed:
                total += parsed[0]
                if parsed[1]:
                    descs.append(parsed[1])
            else:
                descs.append(seg)

        if total > 0:
            results.append((total, " ".join(descs)))

    return results


def append_record(user_id: int, record_type: str, amount: float, payment: str, description: str = "") -> Path:
    path = ensure_workbook(user_id)
    wb = load_workbook(path)
    ws = wb.active

    row = [record_type, amount, description, payment]
    next_row = ws.max_row + 1
    for col_idx, value in enumerate(row, 1):
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.border = THIN_BORDER
        if isinstance(value, float):
            cell.number_format = "#,##0.00"

    wb.save(path)
    return path


def append_records_bulk(user_id: int, record_type: str, entries: list[tuple[float, str]], payment: str) -> Path:
    path = ensure_workbook(user_id)
    wb = load_workbook(path)
    ws = wb.active
    for amount, description in entries:
        row_data = [record_type, amount, description, payment]
        next_row = ws.max_row + 1
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=next_row, column=col_idx, value=value)
            cell.border = THIN_BORDER
            if isinstance(value, float):
                cell.number_format = "#,##0.00"

    wb.save(path)
    return path


async def start(update: Update, context) -> int:
    user_id = update.effective_user.id
    path = get_user_file(user_id)
    if path.exists():
        path.unlink()
    context.user_data.clear()

    keyboard = [["Kirim", "Chiqim"]]
    await update.message.reply_text(
        "Assalomu alaykum! Men sizning kirim-chiqim botingizman.\n\n"
        "Turini tanlang:",
        reply_markup=ReplyKeyboardMarkup(keyboard, one_time_keyboard=True, resize_keyboard=True),
    )
    return TYPE


async def choose_type(update: Update, context) -> int:
    text = update.message.text
    if text not in ("Kirim", "Chiqim"):
        keyboard = [["Kirim", "Chiqim"]]
        await update.message.reply_text(
            "Iltimos, 'Kirim' yoki 'Chiqim' ni tanlang:",
            reply_markup=ReplyKeyboardMarkup(keyboard, one_time_keyboard=True, resize_keyboard=True),
        )
        return TYPE

    context.user_data["type"] = text
    keyboard = [["Yakka", "Ommaviy"]]
    await update.message.reply_text(
        f"{text} tanlandi.\n\nKiritish turini tanlang:",
        reply_markup=ReplyKeyboardMarkup(keyboard, one_time_keyboard=True, resize_keyboard=True),
    )
    return INPUT_MODE


async def choose_input_mode(update: Update, context) -> int:
    text = update.message.text
    if text not in ("Yakka", "Ommaviy"):
        keyboard = [["Yakka", "Ommaviy"]]
        await update.message.reply_text(
            "Iltimos, 'Yakka' yoki 'Ommaviy' ni tanlang:",
            reply_markup=ReplyKeyboardMarkup(keyboard, one_time_keyboard=True, resize_keyboard=True),
        )
        return INPUT_MODE

    context.user_data["input_mode"] = text

    if text == "Yakka":
        await update.message.reply_text(
            "Miqdorni kiriting (masalan: 50000 yoki 10000+25000+3000):",
            reply_markup=ReplyKeyboardRemove(),
        )
        return AMOUNT

    await update.message.reply_text(
        "Ommaviy ma'lumotlarni kiriting.\n"
        "Har bir qator alohida yozuv bo'ladi.\n\n"
        "Misol:\n"
        "110.000 kavchok vilanka 20mtr\n"
        "200.000 sim 4 lik\n"
        "300.000+200.000+500.000",
        reply_markup=ReplyKeyboardRemove(),
    )
    return BULK_TEXT


async def enter_amount(update: Update, context) -> int:
    text = update.message.text.strip()
    result = evaluate_math(text)

    if result is None or result <= 0:
        await update.message.reply_text(
            "Noto'g'ri miqdor. Iltimos, raqam kiriting (masalan: 50000 yoki 10000+25000):"
        )
        return AMOUNT

    context.user_data["amount"] = result

    keyboard = [["Naqd", "Perechisleniya"]]
    display = f"{result:,.2f}".replace(",", " ")
    await update.message.reply_text(
        f"Miqdor: {display} so'm\n\nTo'lov turini tanlang:",
        reply_markup=ReplyKeyboardMarkup(keyboard, one_time_keyboard=True, resize_keyboard=True),
    )
    return PAYMENT


async def enter_bulk(update: Update, context) -> int:
    text = update.message.text.strip()
    entries = parse_bulk_text(text)

    if not entries:
        await update.message.reply_text(
            "Hech qanday yozuv topilmadi. Iltimos, qaytadan kiriting.\n\n"
            "Misol:\n"
            "110.000 kavchok vilanka 20mtr\n"
            "200.000 sim 4 lik\n"
            "300.000+200.000+500.000"
        )
        return BULK_TEXT

    context.user_data["entries"] = entries

    total = sum(a for a, _ in entries)
    display_total = f"{total:,.0f}".replace(",", " ")
    summary_lines = []
    for i, (amount, desc) in enumerate(entries, 1):
        display_amt = f"{amount:,.0f}".replace(",", " ")
        if desc:
            summary_lines.append(f"  {i}. {display_amt} - {desc}")
        else:
            summary_lines.append(f"  {i}. {display_amt}")

    header = f"{len(entries)} ta yozuv topildi:\n"
    footer = f"\n\nJami: {display_total} so'm\n\nTo'lov turini tanlang:"
    budget = 4000 - len(header) - len(footer)

    summary = "\n".join(summary_lines)
    if len(summary) > budget:
        kept = []
        used = 0
        truncated_marker = "  ..."
        for line in summary_lines:
            if used + len(line) + 1 + len(truncated_marker) > budget:
                kept.append(truncated_marker)
                break
            kept.append(line)
            used += len(line) + 1
        summary = "\n".join(kept)

    keyboard = [["Naqd", "Perechisleniya"]]
    await update.message.reply_text(
        f"{header}{summary}{footer}",
        reply_markup=ReplyKeyboardMarkup(keyboard, one_time_keyboard=True, resize_keyboard=True),
    )
    return PAYMENT


async def choose_payment(update: Update, context) -> int:
    text = update.message.text
    if text not in ("Naqd", "Perechisleniya"):
        keyboard = [["Naqd", "Perechisleniya"]]
        await update.message.reply_text(
            "Iltimos, 'Naqd' yoki 'Perechisleniya' ni tanlang:",
            reply_markup=ReplyKeyboardMarkup(keyboard, one_time_keyboard=True, resize_keyboard=True),
        )
        return PAYMENT

    record_type = context.user_data["type"]
    user_id = update.effective_user.id
    input_mode = context.user_data.get("input_mode", "Yakka")

    if input_mode == "Ommaviy":
        entries = context.user_data.get("entries", [])
        path = append_records_bulk(user_id, record_type, entries, text)

        total = sum(a for a, _ in entries)
        display = f"{total:,.0f}".replace(",", " ")
        await update.message.reply_text(
            f"Saqlandi!\n"
            f"Turi: {record_type}\n"
            f"Yozuvlar soni: {len(entries)}\n"
            f"Jami: {display} so'm\n"
            f"To'lov: {text}",
            reply_markup=ReplyKeyboardRemove(),
        )
    else:
        amount = context.user_data["amount"]
        path = append_record(user_id, record_type, amount, text)

        display = f"{amount:,.2f}".replace(",", " ")
        await update.message.reply_text(
            f"Saqlandi!\n"
            f"Turi: {record_type}\n"
            f"Miqdor: {display} so'm\n"
            f"To'lov: {text}",
            reply_markup=ReplyKeyboardRemove(),
        )

    await update.message.reply_text("Fayl nomini kiriting (masalan: may_hisobot):")
    return FILENAME


async def enter_filename(update: Update, context) -> int:
    name = update.message.text.strip()
    if not name:
        await update.message.reply_text("Iltimos, fayl nomini kiriting:")
        return FILENAME

    if not name.endswith(".xlsx"):
        name += ".xlsx"

    user_id = update.effective_user.id
    path = get_user_file(user_id)

    await update.message.reply_document(
        document=open(path, "rb"),
        filename=name,
        caption="Sizning hisobotingiz",
    )

    keyboard = [["Kirim", "Chiqim"]]
    await update.message.reply_text(
        "Yangi yozuv qo'shish uchun turini tanlang yoki /cancel bosing:",
        reply_markup=ReplyKeyboardMarkup(keyboard, one_time_keyboard=True, resize_keyboard=True),
    )
    return TYPE


async def cancel(update: Update, context) -> int:
    await update.message.reply_text(
        "Bekor qilindi. Qaytadan boshlash uchun /start bosing.",
        reply_markup=ReplyKeyboardRemove(),
    )
    return ConversationHandler.END


async def hisobot(update: Update, context) -> None:
    user_id = update.effective_user.id
    path = get_user_file(user_id)
    if not path.exists():
        await update.message.reply_text("Sizda hali yozuvlar yo'q. /start bosing.")
        return
    await update.message.reply_document(
        document=open(path, "rb"),
        filename=f"hisobot_{user_id}.xlsx",
        caption="Sizning hisobotingiz",
    )


def main():
    token = os.environ.get("BOT_TOKEN")
    if not token:
        raise SystemExit("BOT_TOKEN environment variable is not set")

    app = Application.builder().token(token).build()

    conv_handler = ConversationHandler(
        entry_points=[CommandHandler("start", start)],
        states={
            TYPE: [MessageHandler(filters.TEXT & ~filters.COMMAND, choose_type)],
            INPUT_MODE: [MessageHandler(filters.TEXT & ~filters.COMMAND, choose_input_mode)],
            AMOUNT: [MessageHandler(filters.TEXT & ~filters.COMMAND, enter_amount)],
            BULK_TEXT: [MessageHandler(filters.TEXT & ~filters.COMMAND, enter_bulk)],
            PAYMENT: [MessageHandler(filters.TEXT & ~filters.COMMAND, choose_payment)],
            FILENAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, enter_filename)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
    )

    app.add_handler(conv_handler)
    app.add_handler(CommandHandler("hisobot", hisobot))

    print("Bot ishga tushdi...")
    app.run_polling()


if __name__ == "__main__":
    main()
